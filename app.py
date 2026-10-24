import io, os, re, sqlite3, uuid
from datetime import date, datetime, timedelta
from functools import wraps
from flask import (Flask, abort, flash, g, jsonify, redirect, render_template,
                   request, send_file, session, url_for)
from werkzeug.security import check_password_hash, generate_password_hash

BASE = os.path.dirname(os.path.abspath(__file__))
DB = os.path.join(BASE, "receipts.db")
UP = os.path.join(BASE, "static", "uploads")
app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-change-me")
app.permanent_session_lifetime = timedelta(minutes=30)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024
ENTRY = ("admin", "accountant", "staff")
IMG = ("png", "jpg", "jpeg", "gif", "webp")
KINDS = {"course": "courses", "year": "acad_years", "semester": "semesters", "branch": "branches", "purpose": "fee_categories", "mode": "pay_modes"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def db():
    if "db" not in g:
        g.db = sqlite3.connect(DB)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys=ON")
    return g.db


@app.teardown_appcontext
def close_db(_):
    d = g.pop("db", None)
    if d:
        d.close()


def init_db():
    os.makedirs(UP, exist_ok=True)
    con = sqlite3.connect(DB)
    with open(os.path.join(BASE, "schema.sql")) as f:
        con.executescript(f.read())
    if not con.execute("SELECT 1 FROM users").fetchone():
        con.executemany("INSERT INTO users(username,password_hash,role) VALUES(?,?,?)",
                        [("ultra", generate_password_hash("ultra123"), "superadmin"),
                         ("admin", generate_password_hash("admin123"), "admin")])
        for course, dur in (("Diploma", 3), ("BTech", 4)):
            cur = con.execute("INSERT INTO courses(name,duration) VALUES(?,?)", (course, dur))
            con.executemany("INSERT INTO branches(course_id,name) VALUES(?,?)",
                            [(cur.lastrowid, b) for b in ("Electrical", "Mechanical", "Civil", "Computer Science", "Basic Science")])
        con.executemany("INSERT INTO fee_categories(name) VALUES(?)", [(n,) for n in (
            "Admission fee", "Registration fee", "Semester fee", "Examination fee", "Library fee",
            "Laboratory fee", "Hostel fee", "Project fee", "Caution money", "Other")])
        con.executemany("INSERT INTO pay_modes(name) VALUES(?)", [(m,) for m in ("Cash", "UPI", "Bank Transfer", "Cheque")])
        con.execute("INSERT INTO settings(key,value) VALUES('name','GIST, Rayagada')")
    # upgrade older databases
    has = lambda t, col: any(r[1] == col for r in con.execute(f"PRAGMA table_info({t})"))
    if not has("courses", "duration"):
        con.execute("ALTER TABLE courses ADD COLUMN duration INTEGER NOT NULL DEFAULT 4")
        con.execute("UPDATE courses SET duration=3 WHERE lower(name)='diploma'")
    if not has("receipt_items", "semester"):
        con.execute("ALTER TABLE receipt_items ADD COLUMN semester TEXT")
    if not has("users", "email"):
        con.execute("ALTER TABLE users ADD COLUMN email TEXT")
    con.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_users_email ON users(email)")
    if not con.execute("SELECT 1 FROM semesters").fetchone():
        con.executemany("INSERT INTO semesters(name) VALUES(?)", [(str(i),) for i in range(1, 11)])
    # older receipts stored "1st Semester" / "1st Semester Back": turn them into "1" / "1 Back"
    for (v,) in con.execute("SELECT DISTINCT semester FROM receipt_items WHERE semester LIKE '%Semester%'").fetchall():
        m = re.match(r"(\d+)", v)
        if m:
            new = m.group(1) + (" Back" if v.endswith("Back") else "")
            con.execute("UPDATE receipt_items SET semester=? WHERE semester=?", (new, v))
            con.execute("INSERT OR IGNORE INTO semesters(name) VALUES(?)", (new,))
    con.execute("UPDATE students SET acad_session='Unassigned' WHERE acad_session IS NULL OR TRIM(acad_session)=''")
    con.execute("INSERT OR IGNORE INTO acad_years(course_id,label) SELECT DISTINCT c.id, s.acad_session FROM students s JOIN courses c ON c.name=s.course")
    con.commit()
    con.close()


def access(*roles):
    """Ultra admin (superadmin) can open every page."""
    def deco(f):
        @wraps(f)
        def wrapper(*a, **k):
            if "uid" not in session:
                return redirect(url_for("login"))
            if roles and session["role"] != "superadmin" and session["role"] not in roles:
                abort(403)
            return f(*a, **k)
        return wrapper
    return deco


app.jinja_env.globals["has"] = lambda *r: session.get("role") == "superadmin" or session.get("role") in r
app.jinja_env.filters["serial"] = lambda n: "%05d" % n
app.jinja_env.filters["inr"] = lambda x: "₹{:,.2f}".format(x or 0)


@app.context_processor
def inject():
    c = db()
    return dict(college={r["key"]: r["value"] for r in c.execute("SELECT * FROM settings")},
                qrs=c.execute("SELECT * FROM qr_codes ORDER BY id").fetchall())


def modes():
    return [r["name"] for r in db().execute("SELECT name FROM pay_modes ORDER BY id")]


def audit(c, rid, action, old, new):
    c.execute("INSERT INTO audit_log(receipt_id,user_id,action,old_value,new_value,ts) VALUES(?,?,?,?,?,?)",
              (rid, session["uid"], action, old, new, datetime.now().isoformat(timespec="seconds")))


def xlsx_response(wb, name):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=name, mimetype=XLSX)


@app.route('/manifest.json')
def manifest():
    return send_from_directory(os.path.join(app.root_path, 'static'), 'manifest.json')

@app.route('/sw.js')
def service_worker():
    return send_from_directory(os.path.join(app.root_path, 'static'), 'sw.js')

@app.route('/')
def index():
    return "E-Receipt App is running!"

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = db().execute("SELECT * FROM users WHERE (username=? OR email=?) AND active=1",
                       (request.form["username"].strip(), request.form["username"].strip().lower())).fetchone()
        if u and check_password_hash(u["password_hash"], request.form["password"]):
            session.clear()
            session.permanent = True
            session.update(uid=u["id"], user=u["username"], role=u["role"])
            return redirect(url_for("dashboard"))
        flash("Wrong username or password.", "error")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


@app.route("/account", methods=["GET", "POST"])
@access()
def account():
    c, f = db(), request.form
    ultra = session["role"] == "superadmin"
    me = c.execute("SELECT * FROM users WHERE id=?", (session["uid"],)).fetchone()
    if request.method == "POST":
        act = f.get("action")
        try:
            if act == "other":
                if not ultra:
                    abort(403)
                email, new = f.get("email", "").strip().lower(), f.get("new", "")
                if not c.execute("SELECT 1 FROM users WHERE id=?", (f.get("uid"),)).fetchone():
                    flash("Choose a user.", "error")
                elif not email and not new:
                    flash("Enter a new email, a new password, or both.", "error")
                elif email and not EMAIL_RE.match(email):
                    flash("That email address does not look right.", "error")
                elif new and len(new) < 8:
                    flash("The new password must have at least 8 characters.", "error")
                else:
                    if email:
                        c.execute("UPDATE users SET email=? WHERE id=?", (email, f["uid"]))
                    if new:
                        c.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new), f["uid"]))
                    c.commit()
                    flash("That user's login details were updated.")
            elif not check_password_hash(me["password_hash"], f.get("current", "")):
                flash("Your current password is wrong.", "error")
            elif act == "email":
                email = f.get("email", "").strip().lower()
                if not EMAIL_RE.match(email):
                    flash("That email address does not look right.", "error")
                else:
                    c.execute("UPDATE users SET email=? WHERE id=?", (email, me["id"]))
                    c.commit()
                    flash("Your sign-in email is now %s." % email)
            elif act == "password":
                new, again = f.get("new", ""), f.get("again", "")
                if len(new) < 8:
                    flash("The new password must have at least 8 characters.", "error")
                elif new != again:
                    flash("The two new passwords do not match.", "error")
                else:
                    c.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new), me["id"]))
                    c.commit()
                    flash("Your password has been changed.")
        except sqlite3.IntegrityError:
            c.rollback()
            flash("That email address is already used by another account.", "error")
        return redirect(url_for("account"))
    users = c.execute("SELECT id, username, role, email FROM users WHERE active=1 AND id<>? ORDER BY username", (me["id"],)).fetchall() if ultra else []
    return render_template("account.html", me=me, users=users)


@app.route("/")
@access()
def dashboard():
    c, today = db(), date.today().isoformat()
    one = lambda q, *a: c.execute(q, a).fetchone()[0]
    money = "SELECT COALESCE(SUM(amount),0) FROM receipts WHERE status='Issued' AND "
    stats = [
        ("Receipt books", one("SELECT COUNT(*) FROM receipt_books"), "mint"),
        ("Receipts issued", one("SELECT COUNT(*) FROM receipts WHERE status='Issued'"), "mint"),
        ("Today's receipts", one("SELECT COUNT(*) FROM receipts WHERE status='Issued' AND receipt_date=?", today), "amber"),
        ("Today's collection", app.jinja_env.filters["inr"](one(money + "receipt_date=?", today)), "amber"),
        ("This month's collection", app.jinja_env.filters["inr"](one(money + "substr(receipt_date,1,7)=?", today[:7])), "mint"),
        ("Cancelled receipts", one("SELECT COUNT(*) FROM receipts WHERE status='Cancelled'"), "mint"),
    ]
    return render_template("dashboard.html", stats=stats)


def year_label(start, duration):
    return "%d-%d" % (start, start + duration)


def in_use(kind, i):
    q = {"year": "SELECT 1 FROM acad_years y JOIN courses c ON c.id=y.course_id "
                 "JOIN students s ON s.course=c.name AND s.acad_session=y.label WHERE y.id=?",
         "branch": "SELECT 1 FROM branches b JOIN courses c ON c.id=b.course_id "
                   "JOIN students s ON s.course=c.name AND s.branch=b.name WHERE b.id=?",
         "course": "SELECT 1 FROM courses c JOIN students s ON s.course=c.name WHERE c.id=?",
         "semester": "SELECT 1 FROM semesters m JOIN receipt_items i ON i.semester=m.name WHERE m.id=?"}.get(kind)
    return bool(q and db().execute(q, (i,)).fetchone())


def pick(cid, yid=None, bid=None):
    c = db()
    course = c.execute("SELECT * FROM courses WHERE id=?", (cid,)).fetchone() or abort(404)
    year = branch = None
    if yid is not None:
        year = c.execute("SELECT * FROM acad_years WHERE id=? AND course_id=?", (yid, cid)).fetchone() or abort(404)
    if bid is not None:
        branch = c.execute("SELECT * FROM branches WHERE id=? AND course_id=?", (bid, cid)).fetchone() or abort(404)
    return course, year, branch


def cell(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return "" if v is None else str(v).strip()


def group_url(s):
    r = db().execute("SELECT c.id AS cid, y.id AS yid, b.id AS bid FROM courses c "
                     "JOIN acad_years y ON y.course_id=c.id AND y.label=? "
                     "JOIN branches b ON b.course_id=c.id AND b.name=? WHERE c.name=?",
                     (s["acad_session"], s["branch"], s["course"])).fetchone()
    return url_for("st_list", **dict(r)) if r else url_for("students")


# Step 1: choose the course (also a quick search across all students)
@app.route("/students")
@access(*ENTRY)
def students():
    c, q = db(), request.args.get("q", "").strip()
    courses = c.execute("SELECT c.*, (SELECT COUNT(*) FROM students s WHERE s.course=c.name) AS n FROM courses c ORDER BY c.id").fetchall()
    rows = c.execute("SELECT * FROM students WHERE roll_no LIKE ? OR name LIKE ? ORDER BY course, acad_session, branch, id LIMIT 200",
                     (f"%{q}%", f"%{q}%")).fetchall() if q else []
    return render_template("students.html", courses=courses, q=q, rows=rows)


# Step 2: choose the academic year (ultra admin can add / delete years)
@app.route("/students/course/<int:cid>", methods=["GET", "POST"])
@access(*ENTRY)
def st_years(cid):
    c = db()
    course, _, _ = pick(cid)
    if request.method == "POST":
        if session["role"] != "superadmin":
            abort(403)
        f = request.form
        if f.get("action") == "add":
            try:
                s = int(f["start"])
                if not 1990 <= s <= 2100:
                    raise ValueError
                c.execute("INSERT INTO acad_years(course_id,label) VALUES(?,?)", (cid, year_label(s, course["duration"])))
                c.commit()
                flash("Academic year %s added." % year_label(s, course["duration"]))
            except (ValueError, KeyError):
                flash("Enter a valid starting year, for example 2023.", "error")
            except sqlite3.IntegrityError:
                flash("That academic year already exists.", "error")
        elif f.get("action") == "del":
            if in_use("year", f.get("id")):
                flash("This academic year has students. Move or remove them first.", "error")
            else:
                c.execute("DELETE FROM acad_years WHERE id=? AND course_id=?", (f.get("id"), cid))
                c.commit()
                flash("Academic year deleted.")
        return redirect(url_for("st_years", cid=cid))
    years = c.execute("SELECT y.*, (SELECT COUNT(*) FROM students s WHERE s.course=? AND s.acad_session=y.label) AS n "
                      "FROM acad_years y WHERE y.course_id=? ORDER BY y.label DESC", (course["name"], cid)).fetchall()
    return render_template("student_years.html", course=course, years=years)


# Step 3: choose the branch (ultra admin can add / delete branches)
@app.route("/students/course/<int:cid>/year/<int:yid>", methods=["GET", "POST"])
@access(*ENTRY)
def st_branches(cid, yid):
    c = db()
    course, year, _ = pick(cid, yid)
    if request.method == "POST":
        if session["role"] != "superadmin":
            abort(403)
        f = request.form
        if f.get("action") == "add" and f.get("name", "").strip():
            try:
                c.execute("INSERT INTO branches(course_id,name) VALUES(?,?)", (cid, f["name"].strip()))
                c.commit()
                flash("Branch added.")
            except sqlite3.IntegrityError:
                flash("That branch already exists.", "error")
        elif f.get("action") == "del":
            if in_use("branch", f.get("id")):
                flash("This branch has students in some academic year. Remove them first.", "error")
            else:
                c.execute("DELETE FROM branches WHERE id=? AND course_id=?", (f.get("id"), cid))
                c.commit()
                flash("Branch deleted.")
        return redirect(url_for("st_branches", cid=cid, yid=yid))
    branches = c.execute("SELECT b.*, (SELECT COUNT(*) FROM students s WHERE s.course=? AND s.acad_session=? AND s.branch=b.name) AS n "
                         "FROM branches b WHERE b.course_id=? ORDER BY b.name", (course["name"], year["label"], cid)).fetchall()
    return render_template("student_branches.html", course=course, year=year, branches=branches)


# Step 4: student list of one course + year + branch, add one student here
@app.route("/students/course/<int:cid>/year/<int:yid>/branch/<int:bid>", methods=["GET", "POST"])
@access(*ENTRY)
def st_list(cid, yid, bid):
    c = db()
    course, year, branch = pick(cid, yid, bid)
    if request.method == "POST":
        f = request.form
        roll, name = f.get("roll_no", "").strip(), f.get("name", "").strip()
        if not roll or not name:
            flash("Enter the roll number and the student name.", "error")
        else:
            try:
                c.execute("INSERT INTO students(roll_no,name,course,branch,acad_session) VALUES(?,?,?,?,?)",
                          (roll, name, course["name"], branch["name"], year["label"]))
                c.commit()
                flash("Student saved.")
            except sqlite3.IntegrityError:
                flash("That roll number already exists.", "error")
        return redirect(url_for("st_list", cid=cid, yid=yid, bid=bid))
    rows = c.execute("SELECT * FROM students WHERE course=? AND acad_session=? AND branch=? ORDER BY id",
                     (course["name"], year["label"], branch["name"])).fetchall()
    return render_template("student_list.html", course=course, year=year, branch=branch, rows=rows)


@app.route("/students/template")
@access(*ENTRY)
def students_template():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Roll No", "Name"])
    ws.append(["D2026001", "Rahul Kumar"])
    return xlsx_response(wb, "students_template.xlsx")


@app.route("/students/course/<int:cid>/year/<int:yid>/branch/<int:bid>/import", methods=["POST"])
@access(*ENTRY)
def students_import(cid, yid, bid):
    from openpyxl import load_workbook
    course, year, branch = pick(cid, yid, bid)
    back = redirect(url_for("st_list", cid=cid, yid=yid, bid=bid))
    file = request.files.get("file")
    if not file or not file.filename.lower().endswith(".xlsx"):
        flash("Choose an .xlsx file.", "error")
        return back
    try:
        rows = list(load_workbook(file, data_only=True).active.iter_rows(values_only=True))
    except Exception:
        flash("Could not read that file. Save it as a normal .xlsx workbook and try again.", "error")
        return back
    first = 1
    if rows and cell(rows[0][0] if rows[0] else None).lower().replace(" ", "").replace(".", "") in ("rollno", "roll", "rollnumber"):
        first = 2  # skip the header row
    c, added, dup, bad = db(), 0, 0, []
    for n, row in enumerate(rows[first - 1:], start=first):
        if not row or not any(v is not None and str(v).strip() for v in row):
            continue
        roll, name = [cell(v) for v in (list(row) + [None, None])[:2]]
        if not (roll and name):
            bad.append(n)
            continue
        cur = c.execute("INSERT OR IGNORE INTO students(roll_no,name,course,branch,acad_session) VALUES(?,?,?,?,?)",
                        (roll, name, course["name"], branch["name"], year["label"]))
        if cur.rowcount:
            added += 1
        else:
            dup += 1
    c.commit()
    msg = f"{added} students added to {course['name']} {year['label']} {branch['name']}."
    if dup:
        msg += f" {dup} skipped because the roll number already exists."
    if bad:
        msg += " Rows with a missing roll number or name: " + ", ".join(map(str, bad[:15])) + "."
    flash(msg, "warn" if (bad or dup) else "message")
    return back


@app.route("/students/<int:sid>/delete", methods=["POST"])
@access("superadmin")
def student_delete(sid):
    c = db()
    s = c.execute("SELECT * FROM students WHERE id=?", (sid,)).fetchone() or abort(404)
    back = group_url(s)
    if c.execute("SELECT 1 FROM receipts WHERE student_id=?", (sid,)).fetchone():
        flash("%s has receipts, so the student cannot be deleted. Edit the details instead." % s["name"], "error")
    else:
        c.execute("DELETE FROM students WHERE id=?", (sid,))
        c.commit()
        flash("Student %s deleted." % s["name"])
    return redirect(back)


# Ultra admin: correct a student's details
@app.route("/students/<int:sid>/edit", methods=["GET", "POST"])
@access("superadmin")
def student_edit(sid):
    c = db()
    s = c.execute("SELECT * FROM students WHERE id=?", (sid,)).fetchone() or abort(404)
    course = c.execute("SELECT * FROM courses WHERE name=?", (s["course"],)).fetchone()
    cid = course["id"] if course else 0
    years = [r["label"] for r in c.execute("SELECT label FROM acad_years WHERE course_id=? ORDER BY label DESC", (cid,))]
    branches = [r["name"] for r in c.execute("SELECT name FROM branches WHERE course_id=? ORDER BY name", (cid,))]
    if request.method == "POST":
        f = request.form
        roll, name, ses, br = (f.get(k, "").strip() for k in ("roll_no", "name", "acad_session", "branch"))
        if not roll or not name or ses not in years or br not in branches:
            flash("Fill in the roll number and name, and choose a valid academic year and branch.", "error")
        else:
            try:
                c.execute("UPDATE students SET roll_no=?, name=?, acad_session=?, branch=? WHERE id=?", (roll, name, ses, br, sid))
                c.commit()
                flash("Student updated.")
                return redirect(group_url({"course": s["course"], "acad_session": ses, "branch": br}))
            except sqlite3.IntegrityError:
                flash("Another student already has that roll number.", "error")
    return render_template("student_edit.html", s=s, years=years, branches=branches, back=group_url(s))


@app.route("/api/student/<roll>")
@access()
def api_student(roll):
    s = db().execute("SELECT name,course,branch,acad_session FROM students WHERE roll_no=?", (roll,)).fetchone()
    return jsonify(dict(s) if s else {})


@app.route("/books", methods=["GET", "POST"])
@access("admin", "accountant")
def books():
    c = db()
    if request.method == "POST":
        f = request.form
        try:
            a, b = int(f["serial_from"]), int(f["serial_to"])
            if a > b:
                raise ValueError
            c.execute("INSERT INTO receipt_books(book_no,serial_from,serial_to,assigned_to,issue_date) VALUES(?,?,?,?,?)",
                      (f["book_no"].strip(), a, b, f["assigned_to"].strip(), f["issue_date"]))
            c.commit()
            flash("Receipt book registered.")
        except (ValueError, KeyError):
            flash("Enter a valid serial range.", "error")
        except sqlite3.IntegrityError:
            flash("That book number is already registered.", "error")
        return redirect(url_for("books"))
    rows = c.execute("SELECT b.*, (SELECT COUNT(*) FROM receipts r WHERE r.book_id=b.id) AS used "
                     "FROM receipt_books b ORDER BY b.id DESC").fetchall()
    return render_template("books.html", rows=rows, today=date.today().isoformat())


BASE_SQL = ("SELECT r.*, b.book_no, s.name, s.roll_no, s.branch, s.course, s.acad_session, "
            "u.username AS entered_by_name, (SELECT group_concat(k.name||CASE WHEN COALESCE(i.semester,'')<>'' THEN ', Sem '||i.semester ELSE '' END||' ('||printf('%.2f',i.amount)||')', '; ') "
            "FROM receipt_items i JOIN fee_categories k ON k.id=i.category_id WHERE i.receipt_id=r.id) AS purposes "
            "FROM receipts r JOIN receipt_books b ON b.id=r.book_id JOIN students s ON s.id=r.student_id "
            "JOIN users u ON u.id=r.entered_by")


def desc(r):
    return "%s-%05d %s %s %s %s" % (r["book_no"], r["serial_no"], r["roll_no"], r["amount"], r["pay_mode"], r["purposes"])


@app.route("/receipts/new", methods=["GET", "POST"])
@app.route("/receipts/<int:rid>/edit", methods=["GET", "POST"])
@access(*ENTRY)
def receipt_save(rid=None):
    c = db()
    if rid and session["role"] != "superadmin":
        abort(403)
    old = c.execute(BASE_SQL + " WHERE r.id=?", (rid,)).fetchone() if rid else None
    if rid and not old:
        abort(404)
    books_ = c.execute("SELECT * FROM receipt_books WHERE status='In use' OR id=? ORDER BY book_no", (old["book_id"] if old else 0,)).fetchall()
    cats = c.execute("SELECT * FROM fee_categories ORDER BY id").fetchall()
    sems = [r["name"] for r in c.execute("SELECT name FROM semesters ORDER BY id")]
    if request.method == "GET":
        f = dict(old) if old else {}
        items = [{"cat": r["category_id"], "sem": r["semester"] or "", "amt": r["amount"]} for r in c.execute(
            "SELECT * FROM receipt_items WHERE receipt_id=? ORDER BY id", (rid,))] if rid else [{"cat": "", "sem": "", "amt": ""}]
    else:
        f = request.form
        items = [{"cat": k, "sem": s, "amt": a} for k, s, a in zip(f.getlist("category_id"), f.getlist("item_semester"), f.getlist("item_amount"))] or [{"cat": "", "sem": "", "amt": ""}]
        stu = c.execute("SELECT id FROM students WHERE roll_no=?", (f.get("roll_no", "").strip(),)).fetchone()
        book = c.execute("SELECT * FROM receipt_books WHERE id=?", (f.get("book_id"),)).fetchone()
        try:
            sn = int(f["serial_no"])
            parsed = [(int(i["cat"]), i["sem"], float(i["amt"])) for i in items if str(i["amt"]).strip()]
        except (ValueError, KeyError):
            sn, parsed = 0, []
        total = sum(a for _, _, a in parsed)
        if not stu:
            flash("Roll number not found. Add the student first.", "error")
        elif not book or not book["serial_from"] <= sn <= book["serial_to"]:
            flash("Serial number is outside this book's range.", "error")
        elif not parsed or any(a <= 0 or s not in sems + [""] for _, s, a in parsed) or f.get("pay_mode") not in modes() or not f.get("receipt_date"):
            flash("Add at least one purpose with an amount, and check the payment mode and date.", "error")
        else:
            ref = f.get("txn_ref", "").strip()
            if ref and c.execute("SELECT 1 FROM receipts WHERE txn_ref=? AND id IS NOT ?", (ref, rid)).fetchone():
                flash("Warning: this transaction reference is already on another receipt.", "warn")
            vals = (book["id"], sn, stu["id"], f["receipt_date"], total, f["pay_mode"], ref, f.get("remarks", "").strip())
            try:
                if rid:
                    c.execute("UPDATE receipts SET book_id=?,serial_no=?,student_id=?,receipt_date=?,amount=?,pay_mode=?,txn_ref=?,remarks=? WHERE id=?", vals + (rid,))
                    c.execute("DELETE FROM receipt_items WHERE receipt_id=?", (rid,))
                else:
                    rid = c.execute("INSERT INTO receipts(book_id,serial_no,student_id,receipt_date,amount,pay_mode,txn_ref,remarks,entry_ts,entered_by) "
                                    "VALUES(?,?,?,?,?,?,?,?,?,?)", vals + (datetime.now().isoformat(timespec="seconds"), session["uid"])).lastrowid
                c.executemany("INSERT INTO receipt_items(receipt_id,category_id,semester,amount) VALUES(?,?,?,?)", [(rid, k, s, a) for k, s, a in parsed])
                new = c.execute(BASE_SQL + " WHERE r.id=?", (rid,)).fetchone()
                audit(c, rid, "edited" if old else "created", desc(old) if old else None, desc(new))
                c.commit()
                flash("Receipt saved.")
                return redirect(url_for("receipt_view", rid=rid))
            except sqlite3.IntegrityError:
                c.rollback()
                rid = old["id"] if old else None
                flash("This book and serial number already has a receipt.", "error")
    return render_template("receipt_form.html", books=books_, cats=cats, modes=modes(), today=date.today().isoformat(),
                           f=f, items=items, edit=bool(old), sems=sems)


def find(args):
    w, p = [], []
    q = args.get("q", "").strip()
    if q:
        w.append("(b.book_no LIKE ? OR CAST(r.serial_no AS TEXT) LIKE ? OR s.name LIKE ? OR s.roll_no LIKE ?)")
        p += [f"%{q}%", f"%{q.lstrip('0') or q}%", f"%{q}%", f"%{q}%"]
    for key, col, op in (("date_from", "r.receipt_date", ">="), ("date_to", "r.receipt_date", "<="),
                         ("mode", "r.pay_mode", "="), ("status", "r.status", "=")):
        if args.get(key):
            w.append(f"{col}{op}?")
            p.append(args[key])
    for key, col in (("course", "s.course"), ("year", "s.acad_session"), ("branch", "s.branch")):
        if args.get(key):
            w.append(f"{col}=?")
            p.append(args[key])
    if args.get("category") or args.get("semester"):
        cond, cp = "", []
        if args.get("category"):
            cond += " AND i.category_id=?"
            cp.append(args["category"])
        if args.get("semester"):
            cond += " AND i.semester=?"
            cp.append(args["semester"])
        w.append("EXISTS(SELECT 1 FROM receipt_items i WHERE i.receipt_id=r.id" + cond + ")")
        p += cp
    sql = BASE_SQL + (" WHERE " + " AND ".join(w) if w else "") + " ORDER BY r.receipt_date DESC, r.id DESC LIMIT 500"
    return db().execute(sql, p).fetchall()


def paid_summary(args):
    """How many students of a course / academic year have paid a purpose, branch by branch."""
    course, cat, year, sel = args.get("course", ""), args.get("category", ""), args.get("year", ""), args.get("branch", "")
    sem = args.get("semester", "")
    if not (course and cat):
        return None
    c = db()
    purpose = c.execute("SELECT name FROM fee_categories WHERE id=?", (cat,)).fetchone()
    cr = c.execute("SELECT id FROM courses WHERE name=?", (course,)).fetchone()
    if not purpose or not cr:
        return None
    yw, yp = (" AND s.acad_session=?", [year]) if year else ("", [])
    sw, sp = (" AND i.semester=?", [sem]) if sem else ("", [])
    paid = {r["branch"]: r for r in c.execute(
        "SELECT s.branch, COUNT(DISTINCT s.id) AS n, SUM(i.amount) AS amt FROM receipt_items i "
        "JOIN receipts r ON r.id=i.receipt_id AND r.status='Issued' JOIN students s ON s.id=r.student_id "
        "WHERE i.category_id=? AND s.course=?" + yw + sw + " GROUP BY s.branch", [cat, course] + yp + sp)}
    total = {r["branch"]: r["n"] for r in c.execute(
        "SELECT s.branch, COUNT(*) AS n FROM students s WHERE s.course=?" + yw + " GROUP BY s.branch", [course] + yp)}
    names = [r["name"] for r in c.execute("SELECT name FROM branches WHERE course_id=? ORDER BY name", (cr["id"],))]
    names += [b for b in list(paid) + list(total) if b not in names]
    base = args.to_dict()
    rows = [dict(branch=b, paid=paid[b]["n"] if b in paid else 0, total=total.get(b, 0),
                 amount=(paid[b]["amt"] or 0) if b in paid else 0, url=url_for("search", **dict(base, branch=b)),
                 selected=(b == sel)) for b in names]
    chosen = [r for r in rows if r["selected"]]
    shown = chosen or rows
    return dict(purpose=purpose["name"], course=course, year=year, semester=sem, rows=rows, branch=sel,
                paid=sum(r["paid"] for r in shown), total=sum(r["total"] for r in shown), amount=sum(r["amount"] for r in shown),
                all_url=url_for("search", **{k: v for k, v in base.items() if k != "branch"}))


@app.route("/receipts")
@access()
def search():
    c = db()
    cats = c.execute("SELECT * FROM fee_categories ORDER BY id").fetchall()
    courses = c.execute("SELECT * FROM courses ORDER BY id").fetchall()
    years = {x["name"]: [y["label"] for y in c.execute("SELECT label FROM acad_years WHERE course_id=? ORDER BY label DESC", (x["id"],))] for x in courses}
    branches = {x["name"]: [b["name"] for b in c.execute("SELECT name FROM branches WHERE course_id=? ORDER BY name", (x["id"],))] for x in courses}
    sems = [r["name"] for r in c.execute("SELECT name FROM semesters ORDER BY id")]
    return render_template("search.html", rows=find(request.args), cats=cats, sems=sems, modes=modes(), a=request.args,
                           courses=courses, years=years, branches=branches, summary=paid_summary(request.args))


@app.route("/receipts/export")
@access("admin", "accountant", "principal")
def export():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Receipts"
    ws.append(["Book no", "Serial no", "Receipt date", "Entry time", "Student", "Roll no", "Course", "Branch",
               "Session", "Purposes", "Amount", "Payment mode", "Transaction ref", "Status", "Entered by"])
    for r in find(request.args):
        ws.append([r["book_no"], "%05d" % r["serial_no"], r["receipt_date"], r["entry_ts"], r["name"], r["roll_no"],
                   r["course"], r["branch"], r["acad_session"], r["purposes"], r["amount"],
                   r["pay_mode"], r["txn_ref"], r["status"], r["entered_by_name"]])
    return xlsx_response(wb, "receipts.xlsx")


@app.route("/receipts/<int:rid>")
@access()
def receipt_view(rid):
    c = db()
    r = c.execute(BASE_SQL + " WHERE r.id=?", (rid,)).fetchone() or abort(404)
    items = c.execute("SELECT k.name, i.semester, i.amount FROM receipt_items i JOIN fee_categories k ON k.id=i.category_id "
                      "WHERE i.receipt_id=? ORDER BY i.id", (rid,)).fetchall()
    return render_template("receipt_view.html", r=r, items=items, copy=request.args.get("copy"))


@app.route("/receipts/<int:rid>/cancel", methods=["POST"])
@access("admin", "accountant")
def receipt_cancel(rid):
    reason = request.form.get("reason", "").strip()
    if not reason:
        flash("Enter a reason to cancel this receipt.", "error")
    else:
        c = db()
        c.execute("UPDATE receipts SET status='Cancelled', cancel_reason=? WHERE id=? AND status='Issued'", (reason, rid))
        audit(c, rid, "cancelled", "Issued", reason)
        c.commit()
        flash("Receipt cancelled.")
    return redirect(url_for("receipt_view", rid=rid))


@app.route("/daily")
@access()
def daily():
    d, c = request.args.get("d") or date.today().isoformat(), db()
    rows = c.execute("SELECT pay_mode, COUNT(*) AS n, SUM(amount) AS total FROM receipts "
                     "WHERE receipt_date=? AND status='Issued' GROUP BY pay_mode", (d,)).fetchall()
    cancelled = c.execute("SELECT COUNT(*) FROM receipts WHERE receipt_date=? AND status='Cancelled'", (d,)).fetchone()[0]
    return render_template("daily.html", d=d, rows=rows, cancelled=cancelled,
                           n=sum(r["n"] for r in rows), total=sum(r["total"] for r in rows))


@app.route("/college", methods=["GET", "POST"])
@access("admin")
def college_page():
    c, f = db(), request.form
    if request.method == "POST":
        if session["role"] != "superadmin":
            abort(403)
        act = f.get("action")
        if act == "save":
            for k in ("name", "address", "phone", "email", "website",
                      "account_holder", "bank_name", "account_no", "ifsc"):
                c.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (k, f.get(k, "").strip()))
            flash("College details saved.")
        elif act == "qr_add":
            n = 0
            for file in request.files.getlist("images"):
                ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
                if ext in IMG:
                    fn = uuid.uuid4().hex + "." + ext
                    file.save(os.path.join(UP, fn))
                    c.execute("INSERT INTO qr_codes(label,filename) VALUES(?,?)", (f.get("label", "").strip(), fn))
                    n += 1
            flash("%d QR image(s) added." % n if n else "Choose PNG, JPG or WEBP images.", "message" if n else "error")
        elif act == "qr_del":
            row = c.execute("SELECT filename FROM qr_codes WHERE id=?", (f.get("id"),)).fetchone()
            if row:
                c.execute("DELETE FROM qr_codes WHERE id=?", (f["id"],))
                try:
                    os.remove(os.path.join(UP, row["filename"]))
                except OSError:
                    pass
                flash("QR image removed.")
        c.commit()
        return redirect(url_for("college_page"))
    return render_template("college.html")


@app.route("/master", methods=["GET", "POST"])
@access("superadmin")
def master():
    c, f = db(), request.form
    if request.method == "POST" and f.get("kind") == "watermark":
        old = c.execute("SELECT value FROM settings WHERE key='wm_image'").fetchone()

        def drop_old():
            if old and old["value"]:
                try:
                    os.remove(os.path.join(UP, old["value"]))
                except OSError:
                    pass
            c.execute("DELETE FROM settings WHERE key='wm_image'")
        if f.get("action") == "img_del":
            drop_old()
            flash("Watermark image removed.")
        else:
            c.execute("INSERT OR REPLACE INTO settings(key,value) VALUES('wm_text',?)", (f.get("wm_text", "").strip(),))
            file = request.files.get("wm_image")
            if file and file.filename:
                ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
                if ext in IMG:
                    fn = uuid.uuid4().hex + "." + ext
                    file.save(os.path.join(UP, fn))
                    drop_old()
                    c.execute("INSERT OR REPLACE INTO settings(key,value) VALUES('wm_image',?)", (fn,))
                    flash("Watermark saved.")
                else:
                    flash("Text saved, but the image must be PNG, JPG, GIF or WEBP.", "error")
            else:
                flash("Watermark saved.")
        c.commit()
        return redirect(url_for("master"))
    if request.method == "POST":
        kind, table = f.get("kind"), KINDS.get(f.get("kind"))
        if not table:
            abort(400)
        try:
            if f.get("action") == "add":
                if kind == "branch":
                    c.execute("INSERT INTO branches(course_id,name) VALUES(?,?)", (f["course_id"], f["name"].strip()))
                elif kind == "year":
                    dur = c.execute("SELECT duration FROM courses WHERE id=?", (f["course_id"],)).fetchone()["duration"]
                    c.execute("INSERT INTO acad_years(course_id,label) VALUES(?,?)", (f["course_id"], year_label(int(f["start"]), dur)))
                elif kind == "course":
                    c.execute("INSERT INTO courses(name,duration) VALUES(?,?)", (f["name"].strip(), int(f.get("duration") or 4)))
                else:
                    c.execute(f"INSERT INTO {table}(name) VALUES(?)", (f["name"].strip(),))
            elif in_use(kind, f.get("id")):
                flash("This item is in use (by students or receipts), so it cannot be deleted.", "error")
            else:
                c.execute(f"DELETE FROM {table} WHERE id=?", (f["id"],))
            c.commit()
        except (sqlite3.IntegrityError, ValueError, TypeError, KeyError):
            flash("That already exists, is invalid, or is in use and cannot be deleted.", "error")
        return redirect(url_for("master"))
    q = lambda s: c.execute(s).fetchall()
    return render_template("master.html", courses=q("SELECT id, name, name||' ('||duration||' years)' AS label FROM courses ORDER BY id"),
                           years=q("SELECT y.id, c.name||' - '||y.label AS label FROM acad_years y JOIN courses c ON c.id=y.course_id ORDER BY c.id, y.label DESC"),
                           branches=q("SELECT b.id, c.name||' - '||b.name AS label FROM branches b JOIN courses c ON c.id=b.course_id ORDER BY c.id, b.name"),
                           cats=q("SELECT id, name AS label FROM fee_categories ORDER BY id"),
                           sems=q("SELECT id, name AS label FROM semesters ORDER BY id"),
                           pmodes=q("SELECT id, name AS label FROM pay_modes ORDER BY id"))


init_db()
if __name__ == "__main__":
    app.run(debug=True)
